"""Connection group schedule: PDF on LECGI letterhead and Excel workbook."""

import datetime
import os
import tempfile

import pandas as pd
import pymupdf
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import CATEGORIES, CATEGORY_ORDER
from .letterhead import CONTACT, LetterheadDoc, _wrap, logo_png

# Columns of the review sheet. Reviewer-editable ones are highlighted.
CONN_COLS = [
    "conn_id", "page", "member_id", "member", "callout", "role", "end", "category",
    "support_kind", "support", "span_ft", "span_src",
    "V_posted", "Vz_posted", "PT_posted", "PC_posted", "M_posted", "exclude", "notes",
    "mark", "V", "Vz", "PT", "PC", "M", "load_basis", "status",
    "x", "y", "cx", "cy", "text_src",
]
EDITABLE = {"member", "role", "category", "support_kind", "support", "span_ft",
            "V_posted", "Vz_posted", "PT_posted", "PC_posted", "M_posted", "exclude", "notes"}
HEADER_ROW = 6          # table header row on sheets that carry the letterhead block

BRAND_HEX = "800000"
_thin = Side(style="thin", color="999999")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def criteria_lines(cfg):
    return [
        f"Design code: {cfg['code']}.",
        "Loads shown are factored (LRFD) end reactions / forces for connection design.",
        f"Beam shear: load posted on the drawings (not less than {cfg['min_shear_k']:g} k); "
        f"otherwise {cfg['udl_fraction']:.0%} of the beam maximum total uniform load capacity "
        f"Wc = min(8 phiMp / L, 2 phiVn) per AISC 360-22 Ch. F and G, R = {cfg['udl_fraction']:g} Wc / 2, "
        f"minimum {cfg['min_shear_k']:g} k.",
        f"HSS beams: Vy = Vz = {cfg['hss_beam_shear_k']:g} k unless posted.",
        "Brace, truss and moment connection forces must be posted on the drawings or provided by the EOR; "
        "items without them are listed as NEEDS INPUT.",
        "Groups: connections with the same category, member size and support are grouped; the group is "
        "designed for the worst (maximum) load of its members.",
    ]


def notes_lines(cfg):
    return ["Materials: " + "; ".join(cfg["materials"]) + ".",
            f"Bolts: {cfg['bolts']}.  Welds: {cfg['welds']}.",
            cfg["default_connection"],
            "Connection design is NOT part of this schedule. Each group is to be designed in RISA "
            "Connection or IDEA StatiCa for the tabulated loads.",
            "Mark legend: BB beam-to-beam, BC beam-to-column, MBB moment beam-to-beam, "
            "MBC moment beam-to-column, BR bracing, TR truss."]


def _locations(g, limit=60):
    ids = sorted({f"S{r['page']}:{r['member_id'].split('-')[-1]}{r['end']}" for r in g["rows"]})
    s = ", ".join(ids)
    return s if len(s) <= limit else s[:limit].rsplit(",", 1)[0] + f", ... ({len(ids)} total)"


# =====================================================================
#  Connection Summary (structural calculation format)
# =====================================================================
SECTIONS = [("BEAM - BEAM CONNECTIONS", ("BB", "MBB")),
            ("BEAM - COLUMN CONNECTIONS", ("BC", "MBC")),
            ("BRACING CONNECTIONS", ("BR",)),
            ("TRUSS CONNECTIONS", ("TR",))]
SUMMARY_COLS = [   # header, PDF width (pt), Excel width
    ("Conn. #", 34, 9), ("Member Size", 50, 14), ("Member\nConnect to", 56, 16),
    ("*Min\nShear\n(No\nreaction\nshown)\n(kip)", 34, 10), ("*Max\nShear\n(Shown\non dwg)\n(kip)", 34, 10),
    ("* Max\nCapacity\n(kip)", 30, 9), ("Conn.\nRatio\n(Max\nShear/Ca\npacity)", 34, 9),
    ("* Max\nAxial\n(T/C)\n(kip)", 30, 9), ("* Max\nCapacity\nAxial\n(T/C)\n(kip)", 32, 9),
    ("* Max\nmoment\n(Shown\non dwg)\n(kip.ft)", 34, 9), ("* Max\nCapacity\nmoment\n(kip.ft)", 32, 9),
    ("RISA\nConn.\nResult", 30, 9), ("Note", 72, 60),
]


def summary_notes(cfg):
    return ["LRFD is used in this design per General Notes.",
            "Shear loads are obtained from plan and notes.",
            f"Reactions not shown on the drawings: {cfg['udl_fraction']:.0%} of the beam maximum total uniform "
            f"load, not less than {cfg['min_shear_k']:g} kips. HSS beams: Vy=Vz={cfg['hss_beam_shear_k']:g} kips.",
            "Capacity, ratio and RISA result columns are completed from the RISA Connection design."]


def _num(v):
    if v is None or (isinstance(v, float) and v != v):
        return None
    return int(v) if float(v).is_integer() else round(float(v), 1)


def summary_rows(groups):
    """Rows of the Connection Summary: ('section', title) or ('row', dict).
    Conn. # is shown on the first row of each number, member size when it changes."""
    out = []
    for title, cats in SECTIONS:
        gs = [g for g in groups if g["category"] in cats]
        if not gs:
            continue
        out.append(("section", title))
        gs.sort(key=lambda g: (g["mark"][:-2], int(g["mark"][-2:])))     # MBB after BB, stable within
        prev_mark = prev_mem = None
        for g in gs:
            hss_default = g["Vz"] is not None and g["V_dwg"] is None
            note = ""
            if g["status"] == "NEEDS INPUT":
                note = "Loads required from EOR: " + g["load_basis"]
            elif g["status"] == "VERIFY":
                note = "Verify: " + g["load_basis"]
            if g["design_member"] != g["member"]:
                note = "; ".join(x for x in (f"Design member {g['design_member']}", note) if x)
            axial = [v for v in (g["PT"], g["PC"]) if v is not None]
            out.append(("row", dict(
                mark=g["mark"] if g["mark"] != prev_mark else "",
                member=g["member"] if (g["member"], g["mark"]) != (prev_mem, prev_mark) else "",
                support=g["support"],
                vmin=f"Vy=Vz={_num(g['Vz'])}" if hss_default else _num(g["V_min"]),
                vdwg=_num(g["V_dwg"]), axial=_num(max(axial)) if axial else None, moment=_num(g["M"]),
                note=note, status=g["status"], first=g["mark"] != prev_mark)))
            prev_mark, prev_mem = g["mark"], g["member"]
    return out


def _calc_date():
    d = datetime.date.today()
    return f"{d.year}-{d.month}-{d.day}"


def cover_fields(meta):
    no = meta.get("project_no") or ""
    return dict(calc_no=meta.get("calc_no") or (f"{no}-0001" if no else ""), project_no=no,
                name=meta.get("project_name") or "", location=meta.get("project_location") or "",
                prepared=meta.get("prepared_by") or "", checked=meta.get("checked_by") or "",
                reviewed=meta.get("reviewed_by") or "", approved=meta.get("approved_by") or "",
                date=_calc_date())


# =====================================================================
#  PDF
# =====================================================================
def _ctext(page, x, y, text, font="helv", size=9, color=(0, 0, 0), underline=False):
    w = pymupdf.get_text_length(text, fontname=font, fontsize=size)
    page.insert_text((x - w / 2, y), text, fontname=font, fontsize=size, color=color)
    if underline:
        page.draw_line((x - w / 2, y + 1.5), (x + w / 2, y + 1.5), width=0.7)


def _pdf_cover(d, meta):
    """Cover sheet in the LECGI structural calculation format."""
    f = cover_fields(meta)
    p = d.plain_page()
    _ctext(p, 306, 34, "Structural Calculation", "hebo", 11.5, underline=True)
    _ctext(p, 306, 48, "Structural Calculation", "hebo", 10)
    p.draw_rect(pymupdf.Rect(51, 73, 561, 636), width=0.8)
    p.draw_line((51, 148), (561, 148), width=0.8)
    d.logo(p, pymupdf.Rect(56, 79, 174, 124))
    _ctext(p, 380, 106, "Structural Calculation", "hebo", 16, underline=True)
    _ctext(p, 380, 120, f"Calc. No.: {f['calc_no']}", "helv", 8.5)
    _ctext(p, 274, 168, "Connection Design", "hebo", 11.5)
    _ctext(p, 274, 180, "for", "hebo", 11.5)
    for k, ln in enumerate(_wrap(f["name"], 300, "hebo", 11.5)):
        _ctext(p, 274, 193 + 12.5 * k, ln, "hebo", 11.5)
    for k, (lab, val) in enumerate((("Project Number:", f["project_no"]), ("Project Location:", f["location"]))):
        p.insert_text((90, 232 + 13 * k), lab, fontname="helv", fontsize=8.5)
        p.insert_text((196, 232 + 13 * k), val, fontname="helv", fontsize=10)
    for k, (lab, val) in enumerate((("Prepared by:", f["prepared"]), ("Checked by:", f["checked"]),
                                    ("Reviewed by:", f["reviewed"]), ("Approved by:", f["approved"]))):
        y = 345 + 20 * k
        p.insert_text((90, y), lab, fontname="helv", fontsize=8.5)
        p.insert_text((196, y), val, fontname="helv", fontsize=9)
        p.draw_line((194, y + 3), (353, y + 3), width=0.8)
    p.insert_text((196, 418), "Licensed Professional Engineer", fontname="helv", fontsize=8.5)
    _ctext(p, 455, 297, "Stamping for Connection Design", "helv", 8)
    _ctext(p, 455, 309, "of Structural Steel Only", "helv", 8)
    p.draw_rect(pymupdf.Rect(378, 315, 538, 435), width=0.6)
    # revision table
    x = (88, 141, 458, 516)
    y0, rh, nrow = 460, 12.5, 12
    p.draw_rect(pymupdf.Rect(x[0], y0, x[3], y0 + rh * nrow), width=0.9)
    for k in range(1, nrow):
        p.draw_line((x[0], y0 + rh * k), (x[3], y0 + rh * k), width=0.6)
    for xx in x[1:3]:
        p.draw_line((xx, y0), (xx, y0 + rh * nrow), width=0.6)
    for k, row in enumerate((("Rev.", "Revision Description", "Date"), ("0", "First Issue", f["date"]))):
        for j, t in enumerate(row):
            _ctext(p, (x[j] + x[j + 1]) / 2, y0 + rh * k + 9.5, t, "helv", 8.5)
    p.insert_text((51, 766), f"Printed: {f['date']}", fontname="helv", fontsize=9)
    _ctext(p, 306, 766, "Structural Calculation", "helv", 9)


def _summary_cells(r):
    t = lambda v: "" if v is None else str(v)
    return [r["mark"], r["member"], r["support"], t(r["vmin"]), t(r["vdwg"]), "", "",
            t(r["axial"]), "", t(r["moment"]), "", "", r["note"]]


def _pdf_summary(d, groups, cfg):
    size, lh, pad = 6.0, 7.2, 1.6
    widths = [c[1] for c in SUMMARY_COLS]
    widths[-1] = d.BODY.width - sum(widths[:-1])
    xs = [d.BODY.x0 + sum(widths[:j]) for j in range(len(widths) + 1)]
    last = len(widths) - 1

    def cell(j, y, text):
        for k, ln in enumerate(_wrap(text, widths[j] - 2 * pad, "helv", size)):
            tw = pymupdf.get_text_length(ln, fontname="helv", fontsize=size)
            tx = xs[j] + pad if j == last else (xs[j] + xs[j + 1] - tw) / 2
            d.page.insert_text((tx, y + pad + size + k * lh - 0.8), ln, fontname="helv", fontsize=size,
                               color=(0.85, 0, 0) if j == last else (0, 0, 0))

    def header():
        _ctext(d.page, 306, d.y + 10, "Structural Calculation", "hebo", 9, underline=True)
        _ctext(d.page, 306, d.y + 20, "Connection Summary", "hebo", 8)
        d.y += 30
        top = d.y
        hh = max(c[0].count("\n") + 1 for c in SUMMARY_COLS) * lh + 2 * pad
        for j, (h, _w, _x) in enumerate(SUMMARY_COLS):
            lines = h.split("\n")
            for k, ln in enumerate(lines):
                tw = pymupdf.get_text_length(ln, fontname="helv", fontsize=size)
                d.page.insert_text(((xs[j] + xs[j + 1] - tw) / 2, d.y + hh - pad - (len(lines) - 1 - k) * lh - 1),
                                   ln, fontname="helv", fontsize=size)
        d.y += hh
        return top

    def next_page(with_header=True):
        d.page.draw_rect(pymupdf.Rect(xs[0], top, xs[-1], d.y), width=0.6)
        d.new_page()
        return header() if with_header else d.y

    d.new_page()
    top = header()
    for kind, r in summary_rows(groups):
        if kind == "section":
            rh = lh + 2 * pad
        else:
            vals = _summary_cells(r)
            rh = max(len(_wrap(v, w - 2 * pad, "helv", size)) for v, w in zip(vals, widths)) * lh + 2 * pad
        if d.y + rh > d.BODY.y1:
            top = next_page()
        if kind == "section":
            d.page.draw_rect(pymupdf.Rect(xs[0], d.y, xs[-1], d.y + rh), color=(0.4, 0.4, 0.4),
                             fill=(0.85, 0.85, 0.85), width=0.4)
            _ctext(d.page, (xs[0] + xs[-1]) / 2, d.y + pad + size, r, "hebo", size)
        else:
            if r["first"]:
                d.page.draw_line((xs[0], d.y), (xs[-1], d.y), color=(0.55, 0.55, 0.55), width=0.4)
            for j, v in enumerate(vals):
                cell(j, d.y, v)
        d.y += rh
    # notes block, as on the calculation template
    d.page.draw_line((xs[0], d.y), (xs[-1], d.y), width=0.6)
    for s in ["Note:"] + [f"{i}. {s}" for i, s in enumerate(summary_notes(cfg), 1)]:
        for ln in _wrap(s, xs[-1] - xs[0] - 2 * pad, "helv", size):
            if d.y + lh + 2 > d.BODY.y1:
                top = next_page(False)
            d.page.insert_text((xs[0] + pad, d.y + size + 1), ln, fontname="helv", fontsize=size)
            d.y += lh
        d.page.draw_line((xs[0], d.y + 1), (xs[-1], d.y + 1), color=(0.55, 0.55, 0.55), width=0.3)
        d.y += 1.5
    d.page.draw_rect(pymupdf.Rect(xs[0], top, xs[-1], d.y), width=0.6)


def pdf_report(path, groups, rows, page_info, cfg, meta):
    """Structural calculation: cover sheet + Connection Summary on the LECGI letterhead."""
    proj = " - ".join(x for x in (meta.get("project_no"), meta.get("project_name")) if x)
    if meta.get("checked_by"):
        proj = " | ".join(x for x in (proj, f"Checked by: {meta['checked_by']}") if x)
    d = LetterheadDoc("CONNECTION DESIGN", proj)
    _pdf_cover(d, meta)
    _pdf_summary(d, groups, cfg)
    d.save(path)


# =====================================================================
#  Excel
# =====================================================================
def _letterhead_block(ws, title, meta, ncols, logo):
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    if logo:
        img = XLImage(logo)
        img.height, img.width = 62, 133
        ws.add_image(img, "A1")
    last = get_column_letter(max(ncols, 6))
    ws.merge_cells(f"C1:{last}1")
    ws["C1"] = title
    ws["C1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["C1"].fill = PatternFill("solid", fgColor=BRAND_HEX)
    ws["C1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(f"C2:{last}2")
    ws["C2"] = "LECGI   |   Louisville, KY  |  Houston, TX"
    ws["C2"].font = Font(bold=True, color="FFFFFF")
    ws["C2"].fill = PatternFill("solid", fgColor=BRAND_HEX)
    ws["C2"].alignment = Alignment(horizontal="right", vertical="center")
    proj = " - ".join(x for x in (meta.get("project_no"), meta.get("project_name")) if x)
    ws.merge_cells(f"C3:{last}3")
    ws["C3"] = f"Project: {proj or '-'}     Drawings: {meta.get('drawing', '-')}     " \
               f"Date: {datetime.date.today():%m/%d/%Y}" \
               + (f"     Checked by: {meta['checked_by']}" if meta.get("checked_by") else "")
    ws["C3"].font = Font(italic=True, size=9)
    ws.oddFooter.left.text = CONTACT
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P of &N"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"


def _table(ws, header, data, widths=None, fills=None, editable=None, start=HEADER_ROW):
    for j, h in enumerate(header, 1):
        c = ws.cell(row=start, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BRAND_HEX)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    edit_fill = PatternFill("solid", fgColor="FFF9DB")
    for i, row in enumerate(data):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and v != v:
                v = None
            c = ws.cell(row=start + 1 + i, column=j, value=v)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 30)
            if fills and fills[i]:
                c.fill = PatternFill("solid", fgColor=fills[i])
            elif editable and header[j - 1] in editable:
                c.fill = edit_fill
    for j, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(h, max(9, min(40, len(h) + 4)))
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(header))}{start + len(data)}"


def _xl_cover(ws, meta, logo):
    """Cover sheet in the LECGI structural calculation format (columns A-J)."""
    f = cover_fields(meta)
    med = Side(style="medium")
    thin = Side(style="thin")
    for c, w in zip("ABCDEFGHIJ", (3, 9, 9, 9, 12, 12, 8, 10, 12, 3)):
        ws.column_dimensions[c].width = w
    if logo:
        img = XLImage(logo)
        img.height, img.width = 52, 112
        ws.add_image(img, "B2")

    def put(ref, text, size=10, bold=False, underline=False, align="center", merge=None):
        if merge:
            ws.merge_cells(merge)
        ws[ref] = text
        ws[ref].font = Font(size=size, bold=bold, underline="single" if underline else None)
        ws[ref].alignment = Alignment(horizontal=align, vertical="center")

    ws.row_dimensions[3].height = 24
    put("E3", "Structural Calculation", 16, True, True, merge="E3:I3")
    put("E4", f"Calc. No.: {f['calc_no']}", 9, merge="E4:I4")
    put("B8", "Connection Design", 12, True, merge="B8:G8")
    put("B9", "for", 12, True, merge="B9:G9")
    put("B10", f["name"], 12, True, merge="B10:G10")
    for r, (lab, val) in ((12, ("Project Number:", f["project_no"])), (13, ("Project Location:", f["location"]))):
        put(f"B{r}", lab, 9, align="left")
        put(f"D{r}", val, 10, align="left")
    put("G16", "Stamping for Connection Design", 8, merge="G16:I16")
    put("G17", "of Structural Steel Only", 8, merge="G17:I17")
    for r in range(18, 25):
        for c in "GHI":
            ws[f"{c}{r}"].border = Border(left=thin if c == "G" else None, right=thin if c == "I" else None,
                                          top=thin if r == 18 else None, bottom=thin if r == 24 else None)
    for r, lab, val in ((18, "Prepared by:", f["prepared"]), (20, "Checked by:", f["checked"]),
                        (22, "Reviewed by:", f["reviewed"]), (24, "Approved by:", f["approved"])):
        put(f"B{r}", lab, 9, align="left")
        put(f"D{r}", val, 9, align="left")
        for c in "DEF":
            ws[f"{c}{r}"].border = Border(bottom=med)
    put("D25", "Licensed Professional Engineer", 9, align="left")
    # revision table
    for k, row in enumerate([("Rev.", "Revision Description", "Date"), ("0", "First Issue", f["date"])]
                            + [("", "", "")] * 10):
        r = 27 + k
        ws.merge_cells(f"C{r}:G{r}")
        for ref, v in ((f"B{r}", row[0]), (f"C{r}", row[1]), (f"H{r}", row[2])):
            ws[ref] = v
            ws[ref].alignment = Alignment(horizontal="center")
            ws[ref].font = Font(size=9)
        for c in "BCDEFGH":
            ws[f"{c}{r}"].border = Border(top=thin, bottom=thin, left=thin if c in "BCH" else None,
                                          right=thin if c in "BGH" else None)
    # outer frame A1:J40 with the logo band closed at row 6
    for r in range(1, 41):
        for c in "ABCDEFGHIJ":
            cell = ws[f"{c}{r}"]
            b = cell.border
            cell.border = Border(left=med if c == "A" else b.left, right=med if c == "J" else b.right,
                                 top=med if r == 1 else b.top, bottom=med if r in (6, 40) else b.bottom)
    ws.oddHeader.center.text = "&B&U&12Structural Calculation\n&U&11Structural Calculation"
    ws.oddFooter.left.text = f"Printed: {f['date']}"
    ws.oddFooter.center.text = "Structural Calculation"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_area = "A1:J40"
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False


def _xl_summary(ws, groups, cfg, meta, logo):
    """Connection Summary in the structural calculation format."""
    ncol = len(SUMMARY_COLS)
    last = get_column_letter(ncol)
    _letterhead_block(ws, "CONNECTION DESIGN", meta, ncol, logo)
    for r, (text, size) in ((4, ("Structural Calculation", 12)), (5, ("Connection Summary", 11))):
        ws.merge_cells(f"A{r}:{last}{r}")
        ws[f"A{r}"] = text
        ws[f"A{r}"].font = Font(bold=True, size=size, underline="single" if r == 4 else None)
        ws[f"A{r}"].alignment = Alignment(horizontal="center")
    thin = Side(style="thin")
    hair = Side(style="hair", color="999999")
    h = HEADER_ROW + 1
    ws.row_dimensions[h].height = 92
    for j, (text, _w, xw) in enumerate(SUMMARY_COLS, 1):
        c = ws.cell(row=h, column=j, value=text)
        c.font = Font(size=9)
        c.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin if j == 1 else None, right=thin if j == ncol else None)
        ws.column_dimensions[get_column_letter(j)].width = xw
    r = h
    for kind, d in summary_rows(groups):
        r += 1
        side = dict(left=thin, right=thin)
        if kind == "section":
            ws.merge_cells(f"A{r}:{last}{r}")
            ws[f"A{r}"] = d
            ws[f"A{r}"].font = Font(bold=True, size=9)
            ws[f"A{r}"].alignment = Alignment(horizontal="center")
            for j in range(1, ncol + 1):
                ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor="D9D9D9")
                ws.cell(row=r, column=j).border = Border(top=thin, bottom=thin, left=thin if j == 1 else None,
                                                         right=thin if j == ncol else None)
            continue
        vals = [d["mark"], d["member"], d["support"], d["vmin"], d["vdwg"], None,
                f'=IF(N(F{r})=0,"",MAX(D{r}:E{r})/F{r})', d["axial"], None, d["moment"], None, None, d["note"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(size=9, color="D90000" if j == ncol else None)
            c.alignment = Alignment(horizontal="left" if j == ncol else "center", vertical="top",
                                    wrap_text=j == ncol)
            c.border = Border(top=hair if d["first"] else None, left=side["left"] if j == 1 else None,
                              right=side["right"] if j == ncol else None)
        ws.cell(row=r, column=7).number_format = "0.00"
    # notes block
    notes = ["Note:"] + [f"{i}. {s}" for i, s in enumerate(summary_notes(cfg), 1)]
    for k, s in enumerate(notes, 1):
        rr = r + k
        ws.merge_cells(f"A{rr}:{last}{rr}")
        ws[f"A{rr}"] = s
        ws[f"A{rr}"].font = Font(size=9)
        for j in range(1, ncol + 1):
            ws.cell(row=rr, column=j).border = Border(top=thin if k == 1 else None, left=thin if j == 1 else None,
                                                      right=thin if j == ncol else None,
                                                      bottom=thin if k == len(notes) else hair)
    # RISA result: Pass green, Fail red
    from openpyxl.formatting.rule import CellIsRule
    rng = f"L{h + 1}:L{r}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pass"'], font=Font(color="00A000")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fail"'], font=Font(color="D90000")))
    ws.page_setup.orientation = "portrait"
    ws.print_title_rows = f"{h}:{h}"
    ws.freeze_panes = ws.cell(row=h + 1, column=1)
    ws.sheet_view.showGridLines = False


def excel_report(path, groups, rows, page_info, cfg, meta):
    wb = Workbook()
    tmp = os.path.join(tempfile.gettempdir(), "lecgi_logo.png")
    logo = logo_png(tmp)

    ws = wb.active
    ws.title = "Cover"
    _xl_cover(ws, meta, logo)
    _xl_summary(wb.create_sheet("Connection Summary"), groups, cfg, meta, logo)

    ws = wb.create_sheet("Group Schedule")
    hdr = ["Mark", "Category", "Member", "Design member", "Support", "Qty", "Vu (k)", "Vz (k)",
           "Pu T (k)", "Pu C (k)", "Mu (k-ft)", "Status", "Load basis", "Sheets", "Locations"]
    _letterhead_block(ws, "CONNECTION GROUP SCHEDULE", meta, len(hdr), logo)
    data, fills = [], []
    for g in groups:
        data.append([g["mark"], CATEGORIES.get(g["category"], g["category"]), g["member"], g["design_member"],
                     g["support"], g["qty"], g["V"], g["Vz"], g["PT"], g["PC"], g["M"], g["status"],
                     g["load_basis"], ", ".join(map(str, g["sheets"])), _locations(g, 400)])
        fills.append("FFD9B3" if g["status"] == "NEEDS INPUT" else ("FFF3BF" if g["status"] == "VERIFY" else None))
    _table(ws, hdr, data, {"Category": 22, "Load basis": 45, "Locations": 60, "Member": 16,
                           "Design member": 16, "Support": 16}, fills)

    ws = wb.create_sheet("Connections")
    _letterhead_block(ws, "CONNECTION LIST (REVIEW - yellow columns are editable)", meta, 12, logo)
    data = [[r.get(c) for c in CONN_COLS] for r in rows]
    fills = ["FFD9B3" if r.get("status") == "NEEDS INPUT" else None for r in rows]
    _table(ws, CONN_COLS, data, {"notes": 40, "load_basis": 45, "callout": 22}, fills, EDITABLE)
    ws["A4"] = ("Edit yellow columns, then re-run 'build'. category: BB/BC/MBB/MBC/BR/TR (blank = auto). "
                "role: BEAM/MOMENT/BRACE/TRUSS. exclude: any text skips the row. Loads in kips / kip-ft, factored.")
    ws["A4"].font = Font(italic=True, size=9, color="800000")

    ws = wb.create_sheet("Needs Input")
    _letterhead_block(ws, "LOADS REQUIRED FROM EOR", meta, 8, logo)
    need = [r for r in rows if r.get("status") == "NEEDS INPUT"]
    _table(ws, ["conn_id", "page", "mark", "member", "support", "category", "required"],
           [[r["conn_id"], r["page"], r.get("mark"), r["member"], r.get("support") or r.get("support_kind"),
             r["category"], r.get("load_basis")] for r in need], {"required": 60})

    ws = wb.create_sheet("Criteria")
    _letterhead_block(ws, "DESIGN CRITERIA & NOTES", meta, 8, logo)
    for i, s in enumerate(criteria_lines(cfg) + [""] + notes_lines(cfg), HEADER_ROW):
        ws.cell(row=i, column=1, value=s)
    i += 2
    ws.cell(row=i, column=1, value="Sheets processed").font = Font(bold=True)
    for p in page_info:
        i += 1
        ws.cell(row=i, column=1, value=(f"Sheet {p['page']}: {p['n_members']} callouts, scale "
                                        f"{p['scale_ppf'] / 72:.4g}\"=1'-0\" ({p['scale_src']}), "
                                        f"OCR={'yes' if p['ocr'] else 'no'}"))
    ws.column_dimensions["A"].width = 150
    wb.save(path)


def read_review(path):
    """Read the Connections sheet back (from an extract or build workbook)."""
    df = pd.read_excel(path, sheet_name="Connections", header=HEADER_ROW - 1, dtype=object)
    df = df[[c for c in df.columns if c in CONN_COLS]]
    df = df[df["conn_id"].notna() | df["member"].notna()]
    return normalize_rows(df.to_dict("records"))


def normalize_rows(records):
    """Clean reviewer-edited records (Excel or web table) into connection rows."""
    rows = []
    for i, rec in enumerate(records):
        r = {k: (None if (isinstance(v, float) and v != v) else v) for k, v in rec.items()}
        for k in ("member", "category", "role", "support_kind", "support", "end"):
            r[k] = str(r[k]).strip().upper() if r.get(k) not in (None, "") else ""
        if not r.get("member"):
            continue
        if not r.get("conn_id"):
            r["conn_id"] = f"USER-{i + 1}"
            r["member_id"] = r.get("member_id") or r["conn_id"]
        r["exclude"] = "" if r.get("exclude") in (None, "", False) else str(r["exclude"])
        r["notes"] = r.get("notes") or ""
        r["role"] = r["role"] or "BEAM"
        r["end"] = r["end"] or "A"
        rows.append(r)
    return rows
